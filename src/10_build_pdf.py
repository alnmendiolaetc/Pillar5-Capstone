"""
Pillar 5 Capstone - Step 6: Convert the report to PDF
Reads the Word document and rewrites it as a PDF.

We do this with ReportLab rather than Microsoft Word because Word is not
installed on this machine, and because a script can be re-run by anyone.

What this file produces:
    Allen_Mendiola_Credit_Card_Fraud_Detection.pdf

Run after 09_build_report.py:
    python src/10_build_pdf.py
"""

import os

from docx import Document
from docx.shared import Pt

from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Image,
                                Table, TableStyle, PageBreak)
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont


# ===== SETTINGS =====
THIS_FOLDER = os.path.dirname(os.path.abspath(__file__))
PROJECT_FOLDER = os.path.dirname(THIS_FOLDER)

DOCX_NAME = "Allen_Mendiola_Credit_Card_Fraud_Detection.docx"
PDF_NAME = "Allen_Mendiola_Credit_Card_Fraud_Detection.pdf"

DOCX_PATH = os.path.join(PROJECT_FOLDER, DOCX_NAME)
PDF_PATH = os.path.join(PROJECT_FOLDER, PDF_NAME)


# ===== STEP 1: Register fonts =====
# We try to use Calibri so the PDF looks like the Word file. If it is missing
# we quietly fall back to the fonts ReportLab always has.
print("=" * 60)
print("STEP 6: BUILDING THE PDF")
print("=" * 60)

WINDOWS_FONTS = "C:\\Windows\\Fonts"
BODY_FONT = "Helvetica"
BOLD_FONT = "Helvetica-Bold"
ITALIC_FONT = "Helvetica-Oblique"

font_files = [
    ("DocFont", "calibri.ttf"),
    ("DocFont-Bold", "calibrib.ttf"),
    ("DocFont-Italic", "calibrii.ttf"),
]

all_fonts_found = True
for font_name, file_name in font_files:
    font_path = os.path.join(WINDOWS_FONTS, file_name)
    if os.path.exists(font_path):
        pdfmetrics.registerFont(TTFont(font_name, font_path))
    else:
        all_fonts_found = False

if all_fonts_found:
    BODY_FONT = "DocFont"
    BOLD_FONT = "DocFont-Bold"
    ITALIC_FONT = "DocFont-Italic"
    print("Using Calibri fonts.")
else:
    print("Calibri not found, using Helvetica instead.")


# ===== STEP 2: Define how each kind of text should look =====
STYLES = {
    "Title": ParagraphStyle("Title", fontName=BOLD_FONT, fontSize=24, leading=29,
                            alignment=TA_CENTER, spaceAfter=10),
    "Subtitle": ParagraphStyle("Subtitle", fontName=BODY_FONT, fontSize=14, leading=18,
                               alignment=TA_CENTER, textColor=colors.HexColor("#444444"),
                               spaceAfter=8),
    "Heading 1": ParagraphStyle("H1", fontName=BOLD_FONT, fontSize=16, leading=20,
                                spaceBefore=16, spaceAfter=8,
                                textColor=colors.HexColor("#1F3864")),
    "Heading 2": ParagraphStyle("H2", fontName=BOLD_FONT, fontSize=13, leading=17,
                                spaceBefore=12, spaceAfter=6,
                                textColor=colors.HexColor("#2E5496")),
    "Heading 3": ParagraphStyle("H3", fontName=BOLD_FONT, fontSize=11.5, leading=15,
                                spaceBefore=9, spaceAfter=4,
                                textColor=colors.HexColor("#333333")),
    "Normal": ParagraphStyle("Body", fontName=BODY_FONT, fontSize=10, leading=14,
                             alignment=TA_LEFT, spaceAfter=6),
    "Centered": ParagraphStyle("Centered", fontName=BODY_FONT, fontSize=11, leading=15,
                               alignment=TA_CENTER, spaceAfter=4),
    "Bullet": ParagraphStyle("Bullet", fontName=BODY_FONT, fontSize=10, leading=14,
                             leftIndent=18, bulletIndent=6, spaceAfter=3),
    "Caption": ParagraphStyle("Caption", fontName=ITALIC_FONT, fontSize=8.5, leading=11,
                              alignment=TA_CENTER, textColor=colors.HexColor("#444444"),
                              spaceAfter=10),
    "Quote": ParagraphStyle("Quote", fontName=ITALIC_FONT, fontSize=9.5, leading=13,
                            leftIndent=22, textColor=colors.HexColor("#333333"),
                            spaceAfter=6),
    "Code": ParagraphStyle("Code", fontName="Courier", fontSize=8.5, leading=11,
                           leftIndent=22, spaceAfter=2),
}


def escape_for_pdf(text):
    """ReportLab reads a little HTML, so the special characters must be escaped."""
    text = text.replace("&", "&amp;")
    text = text.replace("<", "&lt;")
    text = text.replace(">", "&gt;")
    return text


def paragraph_is_caption(paragraph):
    """Captions are the centred italic lines under each chart."""
    if len(paragraph.runs) == 0:
        return False
    return paragraph.runs[0].italic and paragraph.alignment == 1


def paragraph_is_quote(paragraph):
    """Our callout lines are italic and indented."""
    if len(paragraph.runs) == 0:
        return False
    indent = paragraph.paragraph_format.left_indent
    return paragraph.runs[0].italic and indent is not None


def paragraph_is_code(paragraph):
    if len(paragraph.runs) == 0:
        return False
    return paragraph.runs[0].font.name == "Consolas"


def build_rich_text(paragraph):
    """
    Rebuild a line of text, keeping bold and italic parts.
    This is what preserves the 'Label:' bold prefixes from the Word file.
    """
    pieces = []
    for run in paragraph.runs:
        piece = escape_for_pdf(run.text)
        if run.bold:
            piece = "<b>" + piece + "</b>"
        if run.italic:
            piece = "<i>" + piece + "</i>"
        pieces.append(piece)
    return "".join(pieces)


def add_page_number(canvas, document):
    """Write the page number at the bottom of every page."""
    canvas.saveState()
    canvas.setFont(BODY_FONT, 8)
    canvas.setFillColor(colors.HexColor("#666666"))
    canvas.drawCentredString(A4[0] / 2, 0.45 * inch, str(document.page))
    canvas.drawRightString(A4[0] - 0.85 * inch, 0.45 * inch,
                           "Allen Mendiola - Pillar 5 Capstone")
    canvas.restoreState()


# ===== STEP 3: Read the Word file =====
print("Reading:", DOCX_NAME)
source = Document(DOCX_PATH)

# Word keeps paragraphs, tables and pictures in one flowing body. We walk
# through that body in order so the PDF comes out in the same sequence.
story = []
image_counter = 0

# Pull every picture out of the document in order.
images_in_order = []
for shape in source.inline_shapes:
    image_part = shape._inline.graphic.graphicData.pic.blipFill.blip.embed
    image_part = source.part.related_parts[image_part]
    images_in_order.append(image_part)

body = source.element.body
position_in_paragraphs = 0
position_in_tables = 0

for child in body.iterchildren():
    tag = child.tag.split("}")[-1]

    if tag == "p":
        if position_in_paragraphs >= len(source.paragraphs):
            continue

        paragraph = source.paragraphs[position_in_paragraphs]
        position_in_paragraphs = position_in_paragraphs + 1

        # Does this paragraph hold a picture?
        if "graphic" in child.xml and image_counter < len(images_in_order):
            image_part = images_in_order[image_counter]
            image_counter = image_counter + 1

            temp_name = os.path.join(PROJECT_FOLDER, "reports", "figures",
                                     "_pdf_temp_%d.png" % image_counter)
            handle = open(temp_name, "wb")
            handle.write(image_part.blob)
            handle.close()

            picture = Image(temp_name)
            # Scale it so it fits the page width but keeps its shape.
            max_width = 6.0 * inch
            ratio = picture.imageHeight / float(picture.imageWidth)
            picture.drawWidth = max_width
            picture.drawHeight = max_width * ratio

            # Very tall charts get shrunk so they fit on one page.
            max_height = 7.2 * inch
            if picture.drawHeight > max_height:
                picture.drawHeight = max_height
                picture.drawWidth = max_height / ratio

            picture.hAlign = "CENTER"
            story.append(Spacer(1, 6))
            story.append(picture)
            story.append(Spacer(1, 3))
            continue

        text = paragraph.text.strip()

        # A page break appears as an empty paragraph containing a break.
        if "w:br" in child.xml and 'type="page"' in child.xml:
            story.append(PageBreak())
            if text == "":
                continue

        if text == "":
            story.append(Spacer(1, 5))
            continue

        style_name = paragraph.style.name
        rich_text = build_rich_text(paragraph)

        if style_name in ["Heading 1", "Heading 2", "Heading 3"]:
            story.append(Paragraph(escape_for_pdf(text), STYLES[style_name]))
        elif style_name == "List Bullet":
            story.append(Paragraph(rich_text, STYLES["Bullet"], bulletText="•"))
        elif style_name == "List Number":
            story.append(Paragraph(rich_text, STYLES["Bullet"], bulletText="–"))
        elif paragraph_is_code(paragraph):
            story.append(Paragraph(escape_for_pdf(text), STYLES["Code"]))
        elif paragraph_is_caption(paragraph):
            story.append(Paragraph(rich_text, STYLES["Caption"]))
        elif paragraph_is_quote(paragraph):
            story.append(Paragraph(rich_text, STYLES["Quote"]))
        elif paragraph.alignment == 1:
            # Centred text on the title page.
            size = 11
            if len(paragraph.runs) > 0 and paragraph.runs[0].font.size is not None:
                size = paragraph.runs[0].font.size / Pt(1)

            if size >= 20:
                story.append(Paragraph(escape_for_pdf(text), STYLES["Title"]))
            elif size >= 13:
                story.append(Paragraph(escape_for_pdf(text), STYLES["Subtitle"]))
            else:
                story.append(Paragraph(rich_text, STYLES["Centered"]))
        else:
            story.append(Paragraph(rich_text, STYLES["Normal"]))

    elif tag == "tbl":
        if position_in_tables >= len(source.tables):
            continue

        word_table = source.tables[position_in_tables]
        position_in_tables = position_in_tables + 1

        # Copy the table contents across, wrapping each cell in a Paragraph
        # so that long text wraps instead of running off the page.
        rows = []
        for row_number in range(len(word_table.rows)):
            cells = []
            for cell in word_table.rows[row_number].cells:
                cell_text = escape_for_pdf(cell.text.strip())

                if row_number == 0:
                    cell_style = ParagraphStyle("TableHead", fontName=BOLD_FONT,
                                                fontSize=8, leading=10)
                else:
                    cell_style = ParagraphStyle("TableBody", fontName=BODY_FONT,
                                                fontSize=7.5, leading=9.5)

                cells.append(Paragraph(cell_text, cell_style))
            rows.append(cells)

        if len(rows) == 0:
            continue

        available_width = 6.3 * inch
        column_width = available_width / len(rows[0])

        pdf_table = Table(rows, colWidths=[column_width] * len(rows[0]), repeatRows=1)
        pdf_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DCE6F1")),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#999999")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ("LEFTPADDING", (0, 0), (-1, -1), 4),
            ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1),
             [colors.white, colors.HexColor("#F4F7FB")]),
        ]))

        story.append(Spacer(1, 4))
        story.append(pdf_table)
        story.append(Spacer(1, 8))

print("Collected %d blocks and %d images." % (len(story), image_counter))


# ===== STEP 4: Write the PDF =====
print("Writing:", PDF_NAME)

pdf = SimpleDocTemplate(
    PDF_PATH,
    pagesize=A4,
    leftMargin=0.85 * inch,
    rightMargin=0.85 * inch,
    topMargin=0.75 * inch,
    bottomMargin=0.75 * inch,
    title="Detecting Fraudulent Credit Card Transactions",
    author="Allen Mendiola",
)

pdf.build(story, onFirstPage=add_page_number, onLaterPages=add_page_number)


# ===== STEP 5: Tidy up the temporary image files =====
figures_folder = os.path.join(PROJECT_FOLDER, "reports", "figures")
for file_name in os.listdir(figures_folder):
    if file_name.startswith("_pdf_temp_"):
        os.remove(os.path.join(figures_folder, file_name))

size_mb = os.path.getsize(PDF_PATH) / 1024 / 1024
print()
print("Saved PDF: %s (%.1f MB)" % (PDF_PATH, size_mb))
print()
print("Done. Next step: python src/11_build_slides.py")
